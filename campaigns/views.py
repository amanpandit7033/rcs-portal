"""Views for bulk campaign creation, column mapping, review, HTMX progress tracking, and actions."""
from datetime import datetime
from decimal import Decimal
import logging
from django.contrib import messages
from django.core.paginator import Paginator
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from django.views import View

from accounts.mixins import RoleRequiredMixin
from reports.csv_stream import stream_csv_response
from wallet.models import SenderProfile
from wallet.services import InsufficientBalanceError
from .models import Campaign, CampaignContact
from .services import (
    calculate_campaign_estimate,
    cancel_campaign,
    pause_campaign,
    reserve_campaign_funds,
    resume_campaign,
)
from .tasks import parse_campaign_file_task, run_campaign_task
from .utils import preview_file_rows

logger = logging.getLogger(__name__)


class CampaignListView(RoleRequiredMixin, View):
    """List all bulk campaigns scoped to user role with search and status filtering."""

    allowed_roles = ["admin", "reseller", "user"]
    template_name = "campaigns/campaign_list.html"

    def get(self, request):
        qs = Campaign.objects.for_user(request.user).select_related(
            "user", "sender_profile", "template"
        ).order_by("-created_at")

        status = request.GET.get("status")
        q = request.GET.get("q")

        if status:
            qs = qs.filter(status=status)
        if q:
            qs = qs.filter(name__icontains=q)

        paginator = Paginator(qs, 15)
        page_obj = paginator.get_page(request.GET.get("page", 1))

        return render(
            request,
            self.template_name,
            {
                "campaigns": page_obj,
                "selected_status": status or "",
                "search_query": q or "",
                "page_title": "Bulk RCS Campaigns",
            },
        )


class CampaignCreateView(RoleRequiredMixin, View):
    """Directly configure, preview template, parse recipients, reserve funds, and launch campaign."""

    allowed_roles = ["admin", "reseller", "user"]
    template_name = "campaigns/campaign_create.html"

    def get(self, request):
        if request.user.is_admin or request.user.is_superuser:
            admin_profiles = SenderProfile.objects.filter(user=request.user, is_active=True)
            sender_profiles = admin_profiles if admin_profiles.exists() else SenderProfile.objects.filter(is_active=True)
        else:
            sender_profiles = SenderProfile.objects.filter(user=request.user, is_active=True)

        from templates_mgmt.models import Template

        templates = (
            Template.objects.for_user(request.user)
            .filter(status=Template.Status.APPROVED)
            .select_related("sender_profile")
            .order_by("-created_at")
        )

        import json

        templates_data = {}
        for t in templates:
            payload = t.payload or {}
            stand_alone = payload.get("standAlone") or {}
            templates_data[str(t.id)] = {
                "id": t.id,
                "name": t.name,
                "template_type": t.template_type,
                "message_type": t.message_type,
                "sender_profile_id": str(t.sender_profile_id) if t.sender_profile_id else "",
                "sender_profile_name": t.sender_profile.sender_profile_name if t.sender_profile else "",
                "text_message_content": payload.get("textMessageContent", ""),
                "document_url": payload.get("documentUrl") or payload.get("mediaUrl") or "",
                "media_order": payload.get("mediaOrder") or "media_at_top",
                "card_title": stand_alone.get("cardTitle", ""),
                "card_description": stand_alone.get("cardDescription", ""),
                "card_media_url": stand_alone.get("mediaUrl", ""),
                "card_orientation": stand_alone.get("cardOrientation", "VERTICAL"),
                "card_height": stand_alone.get("cardHeight", "MEDIUM_HEIGHT"),
                "carousel_cards": [
                    {
                        "cardTitle": c.get("cardTitle", ""),
                        "cardDescription": c.get("cardDescription", ""),
                        "mediaUrl": c.get("mediaUrl", ""),
                        "suggestions": c.get("suggestions", []),
                    }
                    for c in payload.get("carouselList", [])
                ],
                "carousel_width": payload.get("carouselWidth", "MEDIUM_WIDTH"),
                "carousel_height": payload.get("carouselHeight", "MEDIUM_HEIGHT"),
                "suggestions": payload.get("suggestions") or stand_alone.get("suggestions") or [],
                "placeholders": t.detected_placeholders,
            }

        user_wallet = getattr(request.user, "wallet", None)
        current_balance = user_wallet.balance if user_wallet else Decimal("0.0000")

        from django.db.models import Count
        from .models import ContactGroup

        contact_groups = (
            ContactGroup.objects.for_user(request.user)
            .annotate(contacts_count=Count("contacts"))
            .order_by("-created_at")
        )

        return render(
            request,
            self.template_name,
            {
                "sender_profiles": sender_profiles,
                "templates": templates,
                "templates_json": json.dumps(templates_data),
                "current_balance": current_balance,
                "contact_groups": contact_groups,
                "page_title": "Create & Send Campaign",
            },
        )

    def post(self, request):
        name = request.POST.get("name", "").strip()
        sender_profile_id = request.POST.get("sender_profile")
        message_type = request.POST.get("message_type", "PROMOTIONAL")
        template_id = request.POST.get("template")
        input_mode = request.POST.get("input_mode", "paste")
        recipient_numbers_raw = request.POST.get("recipient_numbers", "").strip()
        uploaded_file = request.FILES.get("file")
        action_type = request.POST.get("action_type", "now")
        scheduled_datetime_str = request.POST.get("scheduled_at", "").strip()

        if not name or not sender_profile_id:
            messages.error(request, "Please provide campaign name and sender profile.")
            return redirect("campaigns:campaign_create")

        import re
        from django.core.files.base import ContentFile
        from .utils import auto_map_placeholders, detect_mobile_column

        pasted_numbers = []
        if recipient_numbers_raw:
            raw_lines = re.split(r"[\r\n,;]+", recipient_numbers_raw)
            pasted_numbers = [line.strip() for line in raw_lines if line.strip()]

        campaign_file = None
        mobile_col = "mobile_number"
        headers = []

        if input_mode == "file" and uploaded_file:
            campaign_file = uploaded_file
        elif pasted_numbers:
            csv_content = "mobile_number\n" + "\n".join(pasted_numbers) + "\n"
            file_name = f"recipients_{timezone.now().strftime('%Y%m%d_%H%M%S')}.csv"
            campaign_file = ContentFile(csv_content.encode("utf-8"), name=file_name)
        elif uploaded_file:
            campaign_file = uploaded_file
        else:
            messages.error(request, "Please enter recipient mobile numbers or choose a recipient file.")
            return redirect("campaigns:campaign_create")

        if request.user.is_admin or request.user.is_superuser:
            sender_profile = get_object_or_404(SenderProfile, id=sender_profile_id, is_active=True)
        else:
            sender_profile = get_object_or_404(SenderProfile, id=sender_profile_id, user=request.user, is_active=True)

        from templates_mgmt.models import Template

        template = None
        if template_id:
            template = get_object_or_404(
                Template.objects.for_user(request.user),
                id=template_id,
            )

        # Auto-detect message_type from approved template or default to PROMOTIONAL
        final_message_type = (template.message_type if template and template.message_type else message_type) or "PROMOTIONAL"

        # Create campaign record
        campaign = Campaign.objects.create(
            user=request.user,
            name=name,
            sender_profile=sender_profile,
            message_type=final_message_type,
            template=template,
            file=campaign_file,
            status=Campaign.Status.PARSING,
        )

        # Inspect headers for file upload
        if input_mode == "file" or (not pasted_numbers and uploaded_file):
            try:
                preview_data = preview_file_rows(campaign.file.path, max_rows=5)
                headers = preview_data.get("headers", [])
                mobile_col = detect_mobile_column(headers)
            except Exception as e:
                logger.warning("Could not preview file headers for campaign %s: %s", campaign.id, e)
                mobile_col = "mobile_number"

        # Build placeholder parameter mapping
        param_map = {}
        if template:
            auto_mapped = auto_map_placeholders(headers if headers else ["mobile_number"], template.detected_placeholders)
            for p in template.detected_placeholders:
                explicit_col = request.POST.get(f"placeholder_{p}", "").strip()
                if explicit_col:
                    param_map[p] = explicit_col
                elif p in auto_mapped:
                    param_map[p] = auto_mapped[p]

        campaign.column_mapping = {
            "mobile_col": mobile_col,
            "param_map": param_map,
        }
        campaign.save(update_fields=["column_mapping"])

        # Directly parse, reserve funds, and queue or schedule
        parse_campaign_file_task(
            str(campaign.id),
            auto_launch=True,
            schedule_datetime=scheduled_datetime_str if action_type == "schedule" else None,
        )

        campaign.refresh_from_db()

        if campaign.status == Campaign.Status.FAILED:
            messages.error(request, campaign.error_message or "Campaign failed to launch.")
        elif campaign.status == Campaign.Status.SCHEDULED:
            messages.success(request, f"Campaign '{campaign.name}' scheduled for {campaign.scheduled_at.strftime('%Y-%m-%d %H:%M')}.")
        else:
            messages.success(request, f"Campaign '{campaign.name}' launched successfully! {campaign.valid_count} recipients queued.")

        return redirect("campaigns:campaign_detail", campaign_id=campaign.id)


class CampaignDetailView(RoleRequiredMixin, View):
    """Step 4: Campaign report dashboard with live progress, metrics, and actions."""

    allowed_roles = ["admin", "reseller", "user"]
    template_name = "campaigns/campaign_detail.html"

    def get(self, request, campaign_id):
        campaign = get_object_or_404(
            Campaign.objects.for_user(request.user).select_related("user", "sender_profile", "template"),
            id=campaign_id,
        )

        # Get recent contacts sample
        contacts_sample = (
            CampaignContact.objects.filter(campaign=campaign)
            .select_related("message")
            .order_by("id")[:20]
        )

        return render(
            request,
            self.template_name,
            {
                "campaign": campaign,
                "contacts_sample": contacts_sample,
                "page_title": f"Campaign Report - {campaign.name}",
            },
        )


class CampaignProgressAjaxView(RoleRequiredMixin, View):
    """HTMX polling partial returning real-time progress bar and metrics."""

    allowed_roles = ["admin", "reseller", "user"]
    template_name = "campaigns/partials/campaign_progress.html"

    def get(self, request, campaign_id):
        campaign = get_object_or_404(Campaign.objects.for_user(request.user), id=campaign_id)
        return render(
            request,
            self.template_name,
            {"campaign": campaign},
        )


class CampaignActionView(RoleRequiredMixin, View):
    """Handle pause, resume, and cancel commands."""

    allowed_roles = ["admin", "reseller", "user"]

    def post(self, request, campaign_id):
        campaign = get_object_or_404(Campaign.objects.for_user(request.user), id=campaign_id)
        action = request.POST.get("action")

        if action == "pause":
            if pause_campaign(campaign):
                messages.warning(request, f"Campaign '{campaign.name}' paused.")
            else:
                messages.error(request, "Campaign cannot be paused in its current state.")

        elif action == "resume":
            if resume_campaign(campaign):
                messages.success(request, f"Campaign '{campaign.name}' resumed.")
            else:
                messages.error(request, "Campaign cannot be resumed in its current state.")

        elif action == "cancel":
            if cancel_campaign(campaign):
                messages.info(request, f"Campaign '{campaign.name}' cancelled. Unused reservation refunded.")
            else:
                messages.error(request, "Campaign cannot be cancelled.")

        return redirect("campaigns:campaign_detail", campaign_id=campaign.id)


class CampaignExportCSVView(RoleRequiredMixin, View):
    """Stream per-number delivery report as CSV attachment without memory load."""

    allowed_roles = ["admin", "reseller", "user"]

    def get(self, request, campaign_id):
        campaign = get_object_or_404(Campaign.objects.for_user(request.user), id=campaign_id)

        header = [
            "Recipient Number",
            "Contact Status",
            "Message ID",
            "Carrier Message ID",
            "Message Status",
            "Dispatched At",
            "Delivered At",
            "Error / Reason",
        ]

        def format_row(c):
            m = c.message
            return [
                c.number,
                c.get_status_display(),
                str(m.id) if m else "",
                m.onextel_message_id if m and m.onextel_message_id else "",
                m.get_status_display() if m else "",
                m.created_at.strftime("%Y-%m-%d %H:%M:%S") if m and m.created_at else "",
                m.delivered_at.strftime("%Y-%m-%d %H:%M:%S") if m and m.delivered_at else "",
                c.error_message or (m.failure_reason if m else ""),
            ]

        qs = CampaignContact.objects.filter(campaign=campaign).select_related("message").order_by("id")
        filename = f"campaign_{campaign.name.lower().replace(' ', '_')}_results.csv"

        return stream_csv_response(
            filename=filename,
            header=header,
            queryset_iterator=qs.iterator(chunk_size=1000),
            row_formatter=format_row,
        )


# =====================================================================
# Contact Groups Management Views
# =====================================================================

def parse_group_contacts(raw_text: str = "", file_obj=None) -> list:
    """
    Parses contacts from raw text (comma/newline separated) or uploaded file (CSV, Excel, TXT).
    Returns a list of unique tuples: (cleaned_mobile_number, contact_name).
    """
    import re
    import csv
    from messaging.services import normalize_and_validate_indian_phone
    from django.core.exceptions import ValidationError

    contacts_map = {}

    def add_number(raw_num: str, raw_name: str = ""):
        if not raw_num:
            return
        try:
            clean_num = normalize_and_validate_indian_phone(raw_num)
        except ValidationError:
            clean_digits = re.sub(r"[^\d]", "", str(raw_num).strip())
            if len(clean_digits) == 12 and clean_digits.startswith("91"):
                clean_digits = clean_digits[2:]
            elif len(clean_digits) == 11 and clean_digits.startswith("0"):
                clean_digits = clean_digits[1:]
            if 10 <= len(clean_digits) <= 15:
                clean_num = clean_digits
            else:
                return

        if clean_num not in contacts_map:
            contacts_map[clean_num] = raw_name.strip()[:100]

    if raw_text:
        lines = re.split(r"[\r\n]+", raw_text)
        for line in lines:
            line = line.strip()
            if not line:
                continue
            parts = re.split(r"[,;\t]+", line)
            if len(parts) >= 2:
                p0, p1 = parts[0].strip(), parts[1].strip()
                p0_digits = re.sub(r"[^\d]", "", p0)
                p1_digits = re.sub(r"[^\d]", "", p1)
                if len(p0_digits) >= 10:
                    add_number(p0, p1)
                elif len(p1_digits) >= 10:
                    add_number(p1, p0)
                else:
                    add_number(p0)
            elif len(parts) == 1:
                add_number(parts[0])

    if file_obj:
        ext = file_obj.name.lower().split(".")[-1]
        if ext in ["xlsx", "xls"]:
            import openpyxl
            try:
                wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
                sheet = wb.active
                for row in sheet.iter_rows(values_only=True):
                    if not row:
                        continue
                    c0 = str(row[0] or "").strip()
                    c1 = str(row[1] or "").strip() if len(row) > 1 else ""
                    c0_digits = re.sub(r"[^\d]", "", c0)
                    c1_digits = re.sub(r"[^\d]", "", c1)
                    if len(c0_digits) >= 10:
                        add_number(c0, c1)
                    elif len(c1_digits) >= 10:
                        add_number(c1, c0)
            except Exception as e:
                logger.error("Error reading excel file in parse_group_contacts: %s", e)
        else:
            try:
                content = file_obj.read().decode("utf-8", errors="ignore")
                csv_reader = csv.reader(content.splitlines())
                for row in csv_reader:
                    if not row:
                        continue
                    c0 = str(row[0] or "").strip()
                    c1 = str(row[1] or "").strip() if len(row) > 1 else ""
                    c0_digits = re.sub(r"[^\d]", "", c0)
                    c1_digits = re.sub(r"[^\d]", "", c1)
                    if len(c0_digits) >= 10:
                        add_number(c0, c1)
                    elif len(c1_digits) >= 10:
                        add_number(c1, c0)
            except Exception as e:
                logger.error("Error reading csv/txt file in parse_group_contacts: %s", e)

    return [(num, name) for num, name in contacts_map.items()]


class ContactGroupListView(RoleRequiredMixin, View):
    """List contact groups for current user with contact count and search."""

    allowed_roles = ["admin", "reseller", "user"]
    template_name = "campaigns/group_list.html"

    def get(self, request):
        from django.db.models import Count
        from .models import ContactGroup

        q = request.GET.get("q", "").strip()
        qs = (
            ContactGroup.objects.for_user(request.user)
            .annotate(contacts_count=Count("contacts"))
            .order_by("-created_at")
        )
        if q:
            qs = qs.filter(name__icontains=q)

        total_groups = qs.count()
        total_contacts_sum = sum(g.contacts_count for g in qs)

        paginator = Paginator(qs, 15)
        page_obj = paginator.get_page(request.GET.get("page", 1))

        return render(
            request,
            self.template_name,
            {
                "groups": page_obj,
                "search_query": q,
                "total_groups": total_groups,
                "total_contacts_sum": total_contacts_sum,
                "page_title": "Manage Contact Groups",
            },
        )


class ContactGroupCreateView(RoleRequiredMixin, View):
    """Create a new contact group and populate initial contacts."""

    allowed_roles = ["admin", "reseller", "user"]
    template_name = "campaigns/group_form.html"

    def get(self, request):
        return render(
            request,
            self.template_name,
            {
                "page_title": "Create Contact Group",
                "is_create": True,
            },
        )

    def post(self, request):
        from .models import ContactGroup, GroupContact

        name = request.POST.get("name", "").strip()
        description = request.POST.get("description", "").strip()
        raw_numbers = request.POST.get("numbers", "").strip()
        file_obj = request.FILES.get("file")

        if not name:
            messages.error(request, "Please provide a valid Group Name.")
            return redirect("campaigns:group_create")

        group = ContactGroup.objects.create(
            user=request.user,
            name=name,
            description=description,
        )

        parsed_contacts = parse_group_contacts(raw_numbers, file_obj)
        batch = [
            GroupContact(group=group, mobile_number=num, name=contact_name)
            for num, contact_name in parsed_contacts
        ]
        if batch:
            GroupContact.objects.bulk_create(batch, ignore_conflicts=True)

        count = group.contacts.count()
        messages.success(request, f"Group '{name}' created successfully with {count} contacts.")
        return redirect("campaigns:group_detail", pk=group.pk)


class ContactGroupDetailView(RoleRequiredMixin, View):
    """View contacts in a group, search, and add more numbers."""

    allowed_roles = ["admin", "reseller", "user"]
    template_name = "campaigns/group_detail.html"

    def get(self, request, pk):
        from .models import ContactGroup

        group = get_object_or_404(ContactGroup.objects.for_user(request.user), pk=pk)
        q = request.GET.get("q", "").strip()

        contacts_qs = group.contacts.all().order_by("-id")
        if q:
            contacts_qs = contacts_qs.filter(
                Q(mobile_number__icontains=q) | Q(name__icontains=q)
            )

        paginator = Paginator(contacts_qs, 50)
        page_obj = paginator.get_page(request.GET.get("page", 1))

        return render(
            request,
            self.template_name,
            {
                "group": group,
                "contacts": page_obj,
                "search_query": q,
                "total_contacts": group.contacts.count(),
                "page_title": f"Group: {group.name}",
            },
        )

    def post(self, request, pk):
        """Append more contacts to this group."""
        from .models import ContactGroup, GroupContact

        group = get_object_or_404(ContactGroup.objects.for_user(request.user), pk=pk)
        raw_numbers = request.POST.get("numbers", "").strip()
        file_obj = request.FILES.get("file")

        parsed_contacts = parse_group_contacts(raw_numbers, file_obj)
        if not parsed_contacts:
            messages.error(request, "No valid mobile numbers found to add.")
            return redirect("campaigns:group_detail", pk=group.pk)

        initial_count = group.contacts.count()
        batch = [
            GroupContact(group=group, mobile_number=num, name=contact_name)
            for num, contact_name in parsed_contacts
        ]
        GroupContact.objects.bulk_create(batch, ignore_conflicts=True)
        new_count = group.contacts.count() - initial_count

        messages.success(
            request,
            f"Successfully added {new_count} new contacts to '{group.name}' "
            f"({len(parsed_contacts) - new_count} duplicates skipped).",
        )
        return redirect("campaigns:group_detail", pk=group.pk)


class ContactGroupUpdateView(RoleRequiredMixin, View):
    """Edit group name and description."""

    allowed_roles = ["admin", "reseller", "user"]
    template_name = "campaigns/group_edit.html"

    def get(self, request, pk):
        from .models import ContactGroup

        group = get_object_or_404(ContactGroup.objects.for_user(request.user), pk=pk)
        return render(
            request,
            self.template_name,
            {
                "group": group,
                "page_title": f"Edit Group: {group.name}",
            },
        )

    def post(self, request, pk):
        from .models import ContactGroup

        group = get_object_or_404(ContactGroup.objects.for_user(request.user), pk=pk)
        name = request.POST.get("name", "").strip()
        description = request.POST.get("description", "").strip()

        if not name:
            messages.error(request, "Group Name cannot be empty.")
            return redirect("campaigns:group_edit", pk=group.pk)

        group.name = name
        group.description = description
        group.save(update_fields=["name", "description", "updated_at"])

        messages.success(request, f"Group '{name}' updated successfully.")
        return redirect("campaigns:group_detail", pk=group.pk)


class ContactGroupDeleteView(RoleRequiredMixin, View):
    """Delete a contact group and all its members."""

    allowed_roles = ["admin", "reseller", "user"]

    def post(self, request, pk):
        from .models import ContactGroup

        group = get_object_or_404(ContactGroup.objects.for_user(request.user), pk=pk)
        name = group.name
        group.delete()
        messages.success(request, f"Group '{name}' was deleted.")
        return redirect("campaigns:group_list")


class ContactGroupDeleteContactView(RoleRequiredMixin, View):
    """Delete single contact from a group."""

    allowed_roles = ["admin", "reseller", "user"]

    def post(self, request, pk, contact_id):
        from .models import ContactGroup, GroupContact

        group = get_object_or_404(ContactGroup.objects.for_user(request.user), pk=pk)
        contact = get_object_or_404(GroupContact, id=contact_id, group=group)
        contact.delete()
        messages.success(request, f"Contact {contact.mobile_number} removed.")
        return redirect("campaigns:group_detail", pk=group.pk)


class ContactGroupExportView(RoleRequiredMixin, View):
    """Export all contact numbers in a group as CSV."""

    allowed_roles = ["admin", "reseller", "user"]

    def get(self, request, pk):
        from .models import ContactGroup

        group = get_object_or_404(ContactGroup.objects.for_user(request.user), pk=pk)
        qs = group.contacts.all().order_by("id")

        header = ["Mobile Number", "Contact Name", "Added At"]

        def format_row(c):
            return [
                c.mobile_number,
                c.name,
                c.created_at.strftime("%Y-%m-%d %H:%M:%S") if c.created_at else "",
            ]

        safe_name = group.name.lower().replace(" ", "_")
        filename = f"group_{safe_name}_contacts.csv"

        return stream_csv_response(
            filename=filename,
            header=header,
            queryset_iterator=qs.iterator(chunk_size=1000),
            row_formatter=format_row,
        )


class ContactGroupNumbersAjaxView(RoleRequiredMixin, View):
    """API endpoint to fetch all numbers of a group formatted for Campaign textarea insertion."""

    allowed_roles = ["admin", "reseller", "user"]

    def get(self, request, pk):
        from .models import ContactGroup

        group = get_object_or_404(ContactGroup.objects.for_user(request.user), pk=pk)
        numbers = list(group.contacts.values_list("mobile_number", flat=True))

        return JsonResponse(
            {
                "success": True,
                "group_id": group.id,
                "group_name": group.name,
                "count": len(numbers),
                "numbers": numbers,
                "numbers_text": "\n".join(numbers),
            }
        )
