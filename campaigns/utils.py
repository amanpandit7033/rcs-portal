"""File handling and chunk-streaming utilities for CSV and XLSX bulk campaign parsing."""
import csv
import io
import logging
import os
from typing import Any, Dict, Generator, List, Optional, Tuple

logger = logging.getLogger(__name__)


def preview_file_rows(file_path: str, max_rows: int = 10) -> Dict[str, Any]:
    """Read the header and top N rows from CSV or XLSX without loading full file into memory.

    Returns:
        Dict with:
            - headers (List[str]): List of column names.
            - rows (List[List[str]]): List of row values for preview.
            - format (str): 'csv' or 'xlsx'.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Campaign file not found at: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()

    if ext in [".xlsx", ".xlsm", ".xltx", ".xltm"]:
        return _preview_xlsx(file_path, max_rows)
    else:
        return _preview_csv(file_path, max_rows)


def is_likely_header(row) -> bool:
    """Return True if row looks like column names rather than raw data/phone numbers."""
    if not row:
        return True
    cleaned = [str(c or "").strip() for c in row if c is not None and str(c).strip()]
    if not cleaned:
        return True
    # If all items are pure numeric/phone numbers (7+ digits)
    is_all_phone_numbers = all(
        c.replace("+", "").replace("-", "").replace(" ", "").isdigit()
        and len(c.replace("+", "").replace("-", "").replace(" ", "")) >= 7
        for c in cleaned
    )
    return not is_all_phone_numbers


def detect_mobile_column(headers: List[str]) -> str:
    """Auto-detect the phone/mobile column name from headers, or return first column or 'mobile_number'."""
    if not headers:
        return "mobile_number"

    candidates = [
        "mobile", "phone", "mobilenumber", "phonenumber", "mobile_number",
        "phone_number", "contact", "contact_number", "contactnumber",
        "recipient", "number", "dest", "destination", "to", "msisdn",
    ]
    headers_clean = {h.lower().replace("_", "").replace("-", "").replace(" ", ""): h for h in headers}
    for c in candidates:
        clean_c = c.replace("_", "").replace("-", "").replace(" ", "")
        if clean_c in headers_clean:
            return headers_clean[clean_c]

    for h in headers:
        hl = h.lower()
        if any(term in hl for term in ["mob", "phon", "cell", "contact", "msisdn", "num"]):
            return h

    return headers[0]


def auto_map_placeholders(headers: List[str], placeholders: List[str]) -> Dict[str, str]:
    """Auto-match template placeholders to file headers by clean name similarity."""
    if not headers or not placeholders:
        return {}

    headers_clean = {h.lower().replace("_", "").replace("-", "").replace(" ", ""): h for h in headers}
    param_map = {}

    for p in placeholders:
        p_clean = p.lower().replace("_", "").replace("-", "").replace(" ", "")
        if p_clean in headers_clean:
            param_map[p] = headers_clean[p_clean]
        else:
            for h in headers:
                if p.lower() in h.lower():
                    param_map[p] = h
                    break
    return param_map


def _preview_csv(file_path: str, max_rows: int) -> Dict[str, Any]:
    headers: List[str] = []
    rows: List[List[str]] = []

    for enc in ["utf-8-sig", "utf-8", "latin-1"]:
        try:
            with open(file_path, "r", encoding=enc, newline="") as f:
                reader = csv.reader(f)
                first_row = next(reader, None)
                if not first_row:
                    headers = []
                    rows = []
                    break

                if is_likely_header(first_row):
                    headers = [str(col).strip() for col in first_row]
                else:
                    headers = [f"Column {i+1}" for i in range(len(first_row))]
                    rows.append([str(c).strip() for c in first_row])

                for _ in range(max_rows - len(rows)):
                    try:
                        row = next(reader)
                        rows.append([str(c).strip() for c in row])
                    except StopIteration:
                        break
            break
        except UnicodeDecodeError:
            continue

    return {
        "headers": headers,
        "rows": rows,
        "format": "csv",
    }


def _preview_xlsx(file_path: str, max_rows: int) -> Dict[str, Any]:
    import openpyxl

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheet = wb.active

    headers: List[str] = []
    rows: List[List[str]] = []

    try:
        row_iter = sheet.iter_rows(values_only=True)
        first_row = next(row_iter, None)
        if first_row:
            if is_likely_header(first_row):
                headers = [str(col or "").strip() for col in first_row]
            else:
                headers = [f"Column {i+1}" for i in range(len(first_row))]
                rows.append([str(c or "").strip() for c in first_row])

        for _ in range(max_rows - len(rows)):
            row = next(row_iter, None)
            if row is None:
                break
            rows.append([str(c or "").strip() for c in row])
    finally:
        wb.close()

    return {
        "headers": headers,
        "rows": rows,
        "format": "xlsx",
    }


def stream_file_records(
    file_path: str,
    column_mapping: Dict[str, str],
) -> Generator[Tuple[int, str, Dict[str, Any]], None, None]:
    """Generator streaming rows one-by-one from CSV or XLSX."""
    mobile_col = column_mapping.get("mobile_col", "")
    param_map = column_mapping.get("param_map", {})
    ext = os.path.splitext(file_path)[1].lower()

    if ext in [".xlsx", ".xlsm"]:
        yield from _stream_xlsx_records(file_path, mobile_col, param_map)
    else:
        yield from _stream_csv_records(file_path, mobile_col, param_map)


def _stream_csv_records(file_path: str, mobile_col: str, param_map: Dict[str, str]):
    for enc in ["utf-8-sig", "utf-8", "latin-1"]:
        try:
            with open(file_path, "r", encoding=enc, newline="") as f:
                reader = csv.reader(f)
                first_row = next(reader, None)
                if not first_row:
                    return

                has_header = is_likely_header(first_row)
                if has_header:
                    headers = [str(c or "").strip() for c in first_row]
                else:
                    headers = [f"Column {i+1}" for i in range(len(first_row))]

                header_indices = {name: i for i, name in enumerate(headers)}
                if mobile_col not in header_indices and not has_header and len(headers) == 1:
                    mobile_idx = 0
                else:
                    mobile_idx = header_indices.get(mobile_col, 0 if not has_header else None)

                param_indices = {
                    placeholder: header_indices[col]
                    for placeholder, col in param_map.items()
                    if col in header_indices
                }

                data_rows = [first_row] if not has_header else []
                idx = 1
                for row in data_rows:
                    if not row or all(not str(v).strip() for v in row):
                        continue
                    raw_phone = ""
                    if mobile_idx is not None and mobile_idx < len(row):
                        raw_phone = str(row[mobile_idx] or "").strip()
                    params = {
                        placeholder: str(row[p_idx] or "").strip()
                        for placeholder, p_idx in param_indices.items()
                        if p_idx < len(row)
                    }
                    yield idx, raw_phone, params
                    idx += 1

                for row in reader:
                    if not row or all(not str(v).strip() for v in row):
                        continue
                    raw_phone = ""
                    if mobile_idx is not None and mobile_idx < len(row):
                        raw_phone = str(row[mobile_idx] or "").strip()
                    params = {
                        placeholder: str(row[p_idx] or "").strip()
                        for placeholder, p_idx in param_indices.items()
                        if p_idx < len(row)
                    }
                    yield idx, raw_phone, params
                    idx += 1
            break
        except UnicodeDecodeError:
            continue


def _stream_xlsx_records(file_path: str, mobile_col: str, param_map: Dict[str, str]):
    import openpyxl

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheet = wb.active

    try:
        row_iter = sheet.iter_rows(values_only=True)
        first_row = next(row_iter, None)
        if not first_row:
            return

        has_header = is_likely_header(first_row)
        if has_header:
            headers = [str(c or "").strip() for c in first_row]
        else:
            headers = [f"Column {i+1}" for i in range(len(first_row))]

        header_indices = {name: i for i, name in enumerate(headers)}
        if mobile_col not in header_indices and not has_header and len(headers) == 1:
            mobile_idx = 0
        else:
            mobile_idx = header_indices.get(mobile_col, 0 if not has_header else None)

        param_indices = {
            placeholder: header_indices[col]
            for placeholder, col in param_map.items()
            if col in header_indices
        }

        data_rows = [first_row] if not has_header else []
        idx = 1
        for row in data_rows:
            if not row or all(v is None for v in row):
                continue
            raw_phone = ""
            if mobile_idx is not None and mobile_idx < len(row):
                raw_phone = str(row[mobile_idx] or "").strip()
            params = {
                placeholder: str(row[p_idx] or "").strip()
                for placeholder, p_idx in param_indices.items()
                if p_idx < len(row)
            }
            yield idx, raw_phone, params
            idx += 1

        for row in row_iter:
            if not row or all(v is None for v in row):
                continue
            raw_phone = ""
            if mobile_idx is not None and mobile_idx < len(row):
                raw_phone = str(row[mobile_idx] or "").strip()
            params = {
                placeholder: str(row[p_idx] or "").strip()
                for placeholder, p_idx in param_indices.items()
                if p_idx < len(row)
            }
            yield idx, raw_phone, params
            idx += 1

    finally:
        wb.close()
